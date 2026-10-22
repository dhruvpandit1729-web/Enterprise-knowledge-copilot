from __future__ import annotations

import io
from pathlib import Path
import docx
import pytest

from app.ingestion.exceptions import (
    DocumentNotFoundError,
    EmptyDocumentError,
    UnreadableDocumentError,
)
from app.ingestion.parsers.docx import DocxParser


def create_sample_docx(
    elements: list[tuple[str, str] | tuple[str, list[list[str]]]],
) -> bytes:
    """Helper to build a DOCX in memory.

    elements is a list of:
      - ('p', text): paragraph with Normal style
      - ('h1', text): paragraph with Heading 1 style
      - ('h2', text): paragraph with Heading 2 style
      - ('title', text): paragraph with Title style
      - ('tbl', [[cell1, cell2], ...]): table rows
    """
    doc = docx.Document()
    for item_type, data in elements:
        if item_type == "p":
            doc.add_paragraph(data)
        elif item_type == "h1":
            doc.add_heading(data, level=1)
        elif item_type == "h2":
            doc.add_heading(data, level=2)
        elif item_type == "title":
            doc.add_heading(data, level=0)
        elif item_type == "tbl":
            rows = data
            if rows:
                tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for r_idx, row in enumerate(rows):
                    for c_idx, cell in enumerate(row):
                        tbl.cell(r_idx, c_idx).text = cell

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_parser_supported_extensions():
    parser = DocxParser()
    assert parser.supported_extensions == {".docx"}


def test_docx_parser_valid_with_headings_and_tables(tmp_path: Path):
    docx_bytes = create_sample_docx([
        ("p", "Introductory note before any heading."),
        ("h1", "Security Policy"),
        ("p", "All systems must be authenticated."),
        ("h2", "Password Standards"),
        ("p", "Must be at least 12 characters."),
        ("tbl", [["Role", "Min Length"], ["Admin", "16"], ["User", "12"]]),
    ])

    docx_path = tmp_path / "security.docx"
    docx_path.write_bytes(docx_bytes)

    parser = DocxParser()
    doc = parser.parse(docx_path)

    assert doc.source == "security.docx"
    assert doc.metadata["doc_type"] == "docx"
    assert len(doc.sections) == 3

    # Section 0: before headings
    s0 = doc.sections[0]
    assert s0.page_number is None  # Never invent page numbers
    assert s0.heading is None
    assert "Introductory note before any heading." in s0.text

    # Section 1: Security Policy
    s1 = doc.sections[1]
    assert s1.page_number is None
    assert s1.heading == "Security Policy"
    assert "Security Policy" in s1.text
    assert "All systems must be authenticated." in s1.text

    # Section 2: Password Standards (subheading) + table
    s2 = doc.sections[2]
    assert s2.page_number is None
    assert s2.heading == "Security Policy > Password Standards"
    assert "Password Standards" in s2.text
    assert "Must be at least 12 characters." in s2.text
    assert "Admin | 16" in s2.text
    assert "User | 12" in s2.text


def test_docx_parser_no_headings():
    docx_bytes = create_sample_docx([
        ("p", "First paragraph."),
        ("p", "Second paragraph."),
    ])

    parser = DocxParser()
    doc = parser.parse_bytes(docx_bytes, filename="plain.docx")

    assert len(doc.sections) == 1
    assert doc.sections[0].heading is None
    assert doc.sections[0].page_number is None
    assert "First paragraph.\n\nSecond paragraph." in doc.sections[0].text


def test_docx_parser_empty_bytes(tmp_path: Path):
    empty_file = tmp_path / "empty.docx"
    empty_file.write_bytes(b"")

    parser = DocxParser()
    with pytest.raises(EmptyDocumentError, match="empty \\(0 bytes\\)"):
        parser.parse(empty_file)

    with pytest.raises(EmptyDocumentError, match="empty \\(0 bytes\\)"):
        parser.parse_bytes(b"", filename="empty.docx")


def test_docx_parser_blank_document():
    # Empty python-docx document with no paragraphs or tables
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)
    blank_bytes = buf.getvalue()

    parser = DocxParser()
    with pytest.raises(EmptyDocumentError, match="contains no extractable text"):
        parser.parse_bytes(blank_bytes, filename="blank.docx")


def test_docx_parser_corrupt_file(tmp_path: Path):
    corrupt_file = tmp_path / "corrupt.docx"
    corrupt_file.write_bytes(b"PK\x03\x04not-a-valid-docx-zip")

    parser = DocxParser()
    with pytest.raises(UnreadableDocumentError):
        parser.parse(corrupt_file)

    with pytest.raises(UnreadableDocumentError):
        parser.parse_bytes(b"not-a-docx", filename="corrupt.docx")


def test_docx_parser_file_not_found():
    parser = DocxParser()
    with pytest.raises(DocumentNotFoundError):
        parser.parse(Path("missing_file.docx"))
