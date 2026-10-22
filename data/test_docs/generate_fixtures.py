from docx import Document
from docx.shared import Inches
from pathlib import Path

def create_docx_fixture(file_path: Path):
    document = Document()

    document.add_heading('Project Overview', level=1)
    document.add_paragraph('This document outlines the Enterprise Knowledge Copilot project. It is designed to assist users with software engineering tasks by providing relevant information from various internal documents.')

    document.add_heading('Key Components', level=2)
    document.add_paragraph('The system comprises several key components: ingestion, retrieval, and generation.')
    document.add_paragraph('Ingestion handles parsing of different document types like PDF, DOCX, TXT, and Markdown. Retrieval focuses on efficient search within the indexed knowledge base. Generation is responsible for synthesizing answers grounded in evidence.')
    document.add_heading('Technical Specifications', level=2)
    document.add_paragraph('The core development is in Python 3.13.5. We utilize FastAPI for the backend API and Streamlit for the user interface. ChromaDB serves as the vector store for document embeddings.')

    document.add_heading('Testing Standards', level=1)
    document.add_paragraph('All code changes must be accompanied by comprehensive unit and integration tests. Pytest is the chosen framework for test automation. Asynchronous functions require `pytest-asyncio` for proper testing.')
    document.add_paragraph('Before deployment, a full regression suite must pass with zero failures. This ensures stability and reliability across the entire system.')

    document.add_heading('Important Dates', level=2)
    document.add_paragraph('Phase 1 Completion: January 15, 2026')
    document.add_paragraph('Phase 2 Completion: March 30, 2026')
    document.add_paragraph('Final Rollout: June 15, 2026')

    # Add a fact that might be near a chunk boundary
    document.add_paragraph('The maximum supported document size for ingestion is 10MB.')

    document.save(file_path)

if __name__ == "__main__":
    output_path = Path("data/test_docs/multi_section_document.docx")
    create_docx_fixture(output_path)
    print(f"Generated DOCX fixture at {output_path.resolve()}")
